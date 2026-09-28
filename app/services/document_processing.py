import hashlib
import io
import re
import unicodedata
from dataclasses import dataclass

import tiktoken
from docx import Document as WordDocument
from pypdf import PdfReader

from app.core.exceptions import UnsupportedDocumentError, UnsupportedScannedDocumentError


@dataclass(frozen=True)
class PageText:
    page: int | None
    text: str


@dataclass(frozen=True)
class TextChunk:
    page: int | None
    chunk_number: int
    content: str
    content_hash: str


def extract_pages(data: bytes, extension: str) -> list[PageText]:
    normalized_extension = extension.lower().lstrip(".")
    try:
        if normalized_extension == "pdf":
            pages = [
                PageText(page=index, text=_normalize(page.extract_text() or ""))
                for index, page in enumerate(PdfReader(io.BytesIO(data)).pages, start=1)
            ]
            if not any(page.text for page in pages):
                raise UnsupportedScannedDocumentError(
                    "Scanned PDFs without a text layer are not supported."
                )
            return [page for page in pages if page.text]
        if normalized_extension == "docx":
            document = WordDocument(io.BytesIO(data))
            paragraphs = [paragraph.text for paragraph in document.paragraphs if paragraph.text]
            for table in document.tables:
                paragraphs.extend(
                    " | ".join(cell.text.strip() for cell in row.cells)
                    for row in table.rows
                    if any(cell.text.strip() for cell in row.cells)
                )
            text = _normalize("\n".join(paragraphs))
            return [PageText(page=None, text=text)] if text else []
        if normalized_extension in {"txt", "md"}:
            text = _normalize(data.decode("utf-8-sig"))
            return [PageText(page=None, text=text)] if text else []
    except UnsupportedDocumentError:
        raise
    except Exception as exc:
        raise UnsupportedDocumentError("The document could not be parsed.") from exc
    raise UnsupportedDocumentError("Only PDF, DOCX, TXT, and Markdown files are supported.")


def split_pages(
    pages: list[PageText], *, document_id: str, version: str | None, chunk_size: int, overlap: int
) -> list[TextChunk]:
    if chunk_size <= 0 or overlap < 0 or overlap >= chunk_size:
        raise ValueError("Chunk size must be positive and overlap smaller than chunk size.")
    encoding = tiktoken.get_encoding("cl100k_base")
    chunks: list[TextChunk] = []
    chunk_number = 0
    step = chunk_size - overlap
    for page in pages:
        tokens = encoding.encode(page.text)
        for start in range(0, len(tokens), step):
            if start and start >= len(tokens) - overlap:
                break
            content = encoding.decode(tokens[start : start + chunk_size]).strip()
            if content:
                digest_source = f"{document_id}:{version or ''}:{chunk_number}:{content}"
                chunks.append(
                    TextChunk(
                        page=page.page,
                        chunk_number=chunk_number,
                        content=content,
                        content_hash=hashlib.sha256(digest_source.encode("utf-8")).hexdigest(),
                    )
                )
                chunk_number += 1
    return chunks


def _normalize(text: str) -> str:
    normalized = unicodedata.normalize("NFKC", text).replace("\x00", "")
    lines = [re.sub(r"[\t\f\v ]+", " ", line).strip() for line in normalized.splitlines()]
    return re.sub(r"\n{3,}", "\n\n", "\n".join(line for line in lines if line)).strip()
